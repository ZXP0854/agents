#!/usr/bin/env python3
"""生成 Agent 工作原理与执行流程 Word 文档"""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import datetime
import os

doc = Document()

# ========== 样式设置 ==========
style = doc.styles['Normal']
font = style.font
font.name = '宋体'
font.size = Pt(11)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.space_after = Pt(4)

# 标题样式
for i in range(1, 4):
    heading_style = doc.styles[f'Heading {i}']
    heading_style.font.name = '黑体'
    heading_style.font.color.rgb = RGBColor(0x2d, 0x5a, 0x3b)

# ========== 封面 ==========
doc.add_paragraph()
doc.add_paragraph()
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('绿色消费实验 Agent 系统')
run.font.size = Pt(26)
run.font.bold = True
run.font.color.rgb = RGBColor(0x2d, 0x5a, 0x3b)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('工作原理与执行流程说明文档')
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(0x4a, 0x8c, 0x5c)

doc.add_paragraph()
info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
today = datetime.date.today()
info.add_run(f'生成日期：{today.year}年{today.month:02d}月{today.day:02d}日').font.size = Pt(12)

doc.add_paragraph()
doc.add_paragraph()
abstract = doc.add_paragraph()
abstract.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = abstract.add_run('适用于见数（Credamo）实验平台 · 12组被试间设计')
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

doc.add_page_break()

# ========== 一、系统概述 ==========
doc.add_heading('一、系统概述', level=1)

doc.add_paragraph(
    '本系统由 12 个独立的单页 HTML Agent 组成，每个 Agent 对应一个特定的实验条件（产品 × 组别）。'
    'Agent 模拟超市 AI 助手"小林"，在被试输入商品名称后提供对应的环保产品宣传文案，'
    '并能够回答关于产品优缺点、用途、价格、安全性、环保知识等方面的提问。'
)

doc.add_heading('1.1 实验设计', level=2)

doc.add_paragraph(
    '实验采用 2（组别：协作规范诉求组 / 非协作规范诉求组）× 6（产品：A4白纸、一次性纸杯、'
    '金属中性笔、透明胶带大卷、垃圾袋、软毛牙刷）的被试间设计，共 12 个实验条件。'
    '每个被试随机分配至其中一个条件（在见数平台通过问卷分支实现），与该条件对应的专属 Agent 进行交互。'
)

# 表格
table = doc.add_table(rows=13, cols=4)
table.style = 'Light Grid Accent 1'

# 表头
headers = ['序号', '文件名', '实验组别', '对应产品']
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = h
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.bold = True

# 数据
products = ['A4白纸', '一次性纸杯', '金属中性笔', '透明胶带大卷', '垃圾袋', '软毛牙刷']
for idx in range(12):
    row = table.rows[idx + 1]
    group_idx = idx // 6
    prod_idx = idx % 6
    group_label = '协作规范诉求组' if group_idx == 0 else '非协作规范诉求组'

    row.cells[0].text = str(idx + 1).zfill(2)
    row.cells[1].text = f'{idx+1:02d}_{"协作组" if group_idx == 0 else "非协作组"}_{products[prod_idx]}.html'
    row.cells[2].text = group_label
    row.cells[3].text = products[prod_idx]

doc.add_paragraph()

doc.add_heading('1.2 技术架构', level=2)

doc.add_paragraph(
    '每个 Agent HTML 文件均为完全自包含的单文件应用，无外部资源依赖，可直接上传至见数平台并在 iframe 中运行。'
    '技术栈：HTML5 + CSS3 + 原生 JavaScript（ES5 兼容，无框架依赖）。'
)

doc.add_paragraph('核心设计原则：', style='List Bullet')
doc.add_paragraph('文案原文嵌入，确保实验操纵一致性', style='List Bullet 2')
doc.add_paragraph('纯前端运行，零延迟响应用户输入', style='List Bullet 2')
doc.add_paragraph('渐进式引导策略，避免干扰被试自然行为', style='List Bullet 2')
doc.add_paragraph('组别语气自动适配，协作组/非协作组知识回复措辞有差异', style='List Bullet 2')

doc.add_page_break()

# ========== 二、Agent 执行流程 ==========
doc.add_heading('二、Agent 执行流程', level=1)

doc.add_paragraph(
    '以下详细描述被试从加载页面到完成交互的完整流程，以及 Agent 内部的处理逻辑。'
)

doc.add_heading('2.1 页面初始化阶段', level=2)

steps_init = [
    ('步骤 1：页面加载',
     '被试通过见数平台问卷链接进入实验页面。HTML 页面在 iframe 中加载，CSS 渲染对话式界面（绿色简约风格），'
     '顶部显示"绿色消费产品介绍"标题，无组别标签暴露。'),
    ('步骤 2：会话初始化',
     'JavaScript 自动执行以下操作：\n'
     '  (a) 生成唯一会话 ID（session_id），格式为 sess_<timestamp>_<random>\n'
     '  (b) 读取预设的组别变量（MY_GROUP），确定当前 Agent 的协作/非协作归属\n'
     '  (c) 读取预设的产品变量（MY_PRODUCT），确定当前 Agent 负责的产品\n'
     '  (d) 加载产品专属宣传文案（MY_COPY）\n'
     '  (e) 加载产品专属知识库（PRODUCT_KNOWLEDGE，7 维度 × 1 个产品）\n'
     '  (f) 加载通用环保知识库（ECO_KNOWLEDGE，12 条）\n'
     '  (g) 通过 postMessage 向见数父窗口上报 session_init 事件'),
    ('步骤 3：欢迎语发送',
     'Agent 自动发送欢迎消息："顾客您好！欢迎光临小林超市，我是本店的AI智能助手小林，请输入您的商品信息。"'
     '\n注意：欢迎语不透露被试需要输入的具体产品名称，保持超市场景的自然感。'),
]

for title, desc in steps_init:
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.font.bold = True
    doc.add_paragraph(desc)

doc.add_heading('2.2 用户交互处理阶段', level=2)

doc.add_paragraph(
    '被试在输入框中输入文本并点击"发送"（或按 Enter 键）后，Agent 按以下优先级顺序处理输入：'
)

# 流程图描述
doc.add_heading('处理优先级（由高到低）', level=3)

steps_process = [
    ('优先级 1：产品名称匹配',
     'Agent 首先检查用户输入是否匹配当前产品名称或其别名。\n'
     '  匹配规则：\n'
     '  - A4白纸 ← "白纸", "A4纸", "a4纸", "打印纸"\n'
     '  - 一次性纸杯 ← "纸杯", "杯子", "水杯"\n'
     '  - 金属中性笔 ← "中性笔", "笔", "钢笔", "圆珠笔", "签字笔"\n'
     '  - 透明胶带大卷 ← "胶带", "透明胶带"\n'
     '  - 垃圾袋 ← "垃圾袋", "塑料袋", "袋子"\n'
     '  - 软毛牙刷 ← "牙刷", "毛刷"\n'
     '  匹配成功 → 显示对应组别的完整宣传文案（约200字），连续无关输入计数归零。'
     '  文案内容见附件实验参考文本。'),

    ('优先级 2：产品专属知识匹配',
     '若未匹配到产品名称，Agent 检查输入是否包含产品知识关键词：\n'
     '  - 优点类："优点", "优势", "好处", "特色", "亮点"\n'
     '  - 缺点类："缺点", "不足", "劣势", "缺陷"\n'
     '  - 用途类："用途", "怎么用", "使用场景", "适合"\n'
     '  - 对比类："对比", "区别", "比较", "和传统", "vs"\n'
     '  - 价格类："价格", "多少钱", "贵", "成本"\n'
     '  - 安全类："安全", "有毒", "健康", "放心"\n'
     '  - 耐用类："耐用", "质量", "寿命", "持久"\n'
     '  匹配成功 → 返回对应维度的产品专属知识（经组别语气适配处理）。'),

    ('优先级 3：通用环保知识匹配',
     '若前两步均未匹配，Agent 检查输入是否包含环保知识关键词：\n'
     '  "环保", "可降解", "绿色消费", "塑料污染", "碳足迹", "可持续发展",\n'
     '  "竹纤维", "玉米淀粉", "咖啡渣", "甘蔗浆", "PLA", "循环经济"\n'
     '  匹配成功 → 返回对应的通用环保知识（经组别语气适配处理）。'),

    ('优先级 4：无关/错误输入处理',
     '若以上均未匹配，Agent 将输入判定为无关或错误输入，启动渐进式引导：\n'
     '  第 1 次失败："未识别到该商品，请重新输入正确的商品名称。"\n'
     '  第 2 次失败："仍未识别到商品信息，请确认您输入的商品名称是否正确。"\n'
     '  第 3 次及以后："提示：您可以输入【产品名】查询该商品的环保替代信息。'
     '您也可以询问该商品的优缺点、用途、价格等问题。"\n'
     '  设计理念：前两次不提示具体产品名称，保持超市场景的自然交互感；'
     '只有多次失败后才给出明确提示，避免实验者效应。'),

    ('优先级 5：兜底通用响应',
     '若输入不是无关输入但也未被任何知识库匹配（如较长的陈述性内容），'
     'Agent 返回通用的引导消息，列举可询问的内容类型，并保持对话自然延续。'),
]

for title, desc in steps_process:
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.font.bold = True
    doc.add_paragraph(desc)

doc.add_heading('2.3 组别语气适配机制', level=2)

doc.add_paragraph(
    '为体现实验操纵的核心差异（协作规范 vs 非协作规范），Agent 在返回产品知识和环保知识时，'
    '通过 applyGroupTone() 函数对文本进行组别语气后处理：'
)

table_tone = doc.add_table(rows=8, cols=2)
table_tone.style = 'Light Grid Accent 1'
table_tone.rows[0].cells[0].text = '协作规范诉求组'
table_tone.rows[0].cells[1].text = '非协作规范诉求组'
for p in table_tone.rows[0].cells[0].paragraphs:
    for r in p.runs: r.font.bold = True
for p in table_tone.rows[0].cells[1].paragraphs:
    for r in p.runs: r.font.bold = True

tone_data = [
    ('文案开头', '"加入我们，让我们一起为环保做出贡献！"', '"请选择绿色消费！"'),
    ('文案结尾', '"让绿色办公/生活成为共同常态"', '"让绿色办公/生活成为个人习惯"'),
    ('主语称谓', '"我们"（集体视角）', '"您"（个体视角）'),
    ('行动描述', '"共同选择""一起使用"', '"您选择""您使用""个人选择"'),
    ('知识结尾', '"让我们一起为环保做出贡献！"', '"请您通过每一次消费选择，养成绿色生活的个人习惯"'),
    ('消费者指代', '"我们"', '"您"'),
    ('规范诉求', '描述性社会规范 + 共同行动邀请', '描述性社会规范 + 个体行为呼吁'),
]

for i, (dim, collab, noncollab) in enumerate(tone_data):
    row = table_tone.rows[i + 1]
    row.cells[0].text = f'{dim}：{collab}'
    row.cells[1].text = f'{dim}：{noncollab}'

doc.add_paragraph()

doc.add_heading('2.4 数据记录与回调', level=2)

doc.add_paragraph(
    'Agent 全程记录用户交互行为，支持通过以下两种方式回传至见数平台：'
)

doc.add_paragraph('postMessage 回传：每次交互事件（用户输入、系统响应、错误等）通过 window.parent.postMessage() 向见数父窗口发送结构化 JSON 数据。', style='List Bullet')
doc.add_paragraph('本地日志：所有交互记录存储在 interactionLog 数组中，可通过 window.getExperimentData() 方法读取。', style='List Bullet')

doc.add_paragraph('记录的字段包括：')
doc.add_paragraph('sessionId：会话唯一标识', style='List Bullet 2')
doc.add_paragraph('eventType：事件类型（session_init / user_input / product_response / product_knowledge_response / eco_knowledge_response / guidance_response / fallback_response）', style='List Bullet 2')
doc.add_paragraph('product：当前 Agent 负责的产品名称', style='List Bullet 2')
doc.add_paragraph('input：用户输入的原始文本', style='List Bullet 2')
doc.add_paragraph('response：系统返回的文本内容', style='List Bullet 2')
doc.add_paragraph('guidanceLevel：连续无关输入次数（用于衡量被试配合度）', style='List Bullet 2')
doc.add_paragraph('timestamp：ISO 8601 格式时间戳', style='List Bullet 2')

doc.add_page_break()

# ========== 三、技术实现细节 ==========
doc.add_heading('三、技术实现细节', level=1)

doc.add_heading('3.1 文件结构', level=2)

doc.add_paragraph(
    '每个 Agent HTML 文件包含三个内嵌部分：'
)

doc.add_paragraph('HTML 结构层：定义对话界面布局（标题栏、对话区、输入区），使用 Flexbox 实现自适应。', style='List Bullet')
doc.add_paragraph('CSS 样式层：低饱和度绿色系配色（--green-dark: #2d5a3b / --green-mid: #4a8c5c），用户气泡右对齐绿色背景白色文字，系统气泡左对齐浅灰背景深色文字，支持 PC/移动端响应式。', style='List Bullet')
doc.add_paragraph('JavaScript 逻辑层：所有业务逻辑（IIFE 封装，无全局变量污染），包含匹配引擎、知识库、语气适配、数据回调等模块。', style='List Bullet')

doc.add_heading('3.2 模糊匹配算法', level=2)

doc.add_paragraph(
    '产品名称匹配采用两级模糊匹配策略：'
)
doc.add_paragraph('第一级：精确匹配 — 用户输入去除首尾空格后与产品标准名称精确比对', style='List Number')
doc.add_paragraph('第二级：子串匹配 — 遍历产品的别名列表，使用 indexOf 检查用户输入是否包含别名关键词（大小写不敏感）', style='List Number')

doc.add_paragraph(
    '此设计确保"白纸""A4纸""打印纸"等口语化表达均能正确映射至"A4白纸"。'
    '匹配规则遵循"最长精确匹配优先"原则，避免"笔"错误匹配至多个产品。'
)

doc.add_heading('3.3 会话状态管理', level=2)

doc.add_paragraph(
    '会话 ID 在页面加载时通过时间戳（36进制）+ 9 位随机字符串生成，格式为 sess_<ts36>_<random>。'
    '同一页面生命周期内（iframe 未被销毁前）会话 ID 保持不变。'
    '组别和产品变量在页面初始化时从预设常量中读取，运行期间不可变更。'
    '连续无关输入计数器（consecutiveIrrelevant）在每次成功匹配产品后归零。'
)

doc.add_heading('3.4 UI 交互规范', level=2)

doc.add_paragraph('输入框支持 Enter 键快捷发送', style='List Bullet')
doc.add_paragraph('发送后输入框自动清空并聚焦', style='List Bullet')
doc.add_paragraph('系统响应前显示"小林正在为您查询…"加载状态（300ms 模拟延迟）', style='List Bullet')
doc.add_paragraph('新消息自动滚动至可视区域（requestAnimationFrame 优化）', style='List Bullet')
doc.add_paragraph('气泡采用淡入动画（fadeIn 0.25s）', style='List Bullet')
doc.add_paragraph('对话区支持触摸滚动（-webkit-overflow-scrolling: touch）', style='List Bullet')

doc.add_page_break()

# ========== 四、部署说明 ==========
doc.add_heading('四、部署说明', level=1)

doc.add_heading('4.1 见数平台上传步骤', level=2)

deploy_steps = [
    '在见数问卷设计页面，找到对应实验条件的 HTML 模块（每个产品场景中有一个"编辑HTML"按钮）',
    '点击"编辑HTML"，在弹出的编辑器中清空原有内容',
    '打开对应的 Agent HTML 文件（如 01_协作组_A4白纸.html），全选复制全部代码',
    '粘贴到编辑器文本框中，点击"确定"保存',
    '重复以上步骤，完成全部 12 个 HTML 模块的替换',
    '使用见数的"模拟作答"功能验证每个 Agent 的交互行为是否正确',
]

for i, step in enumerate(deploy_steps, 1):
    doc.add_paragraph(f'{i}. {step}')

doc.add_heading('4.2 文件清单', level=2)

doc.add_paragraph('12 个 Agent HTML 文件，命名格式为 <序号>_<组别>_<产品>.html：')
for idx in range(12):
    group_idx = idx // 6
    prod_idx = idx % 6
    group_label = '协作组' if group_idx == 0 else '非协作组'
    doc.add_paragraph(f'{idx+1:02d}_{group_label}_{products[prod_idx]}.html', style='List Bullet')

doc.add_paragraph()
doc.add_paragraph('附带文件：')
doc.add_paragraph('generate_agents.py — Python 生成器脚本，用于重新生成全部 12 个文件', style='List Bullet')
doc.add_paragraph('experiment.html — 通用版（保留 URL 参数方式，支持 ?group=collaboration / non_collaboration）', style='List Bullet')

doc.add_heading('4.3 兼容性说明', level=2)

doc.add_paragraph('浏览器支持：Chrome（推荐）、Microsoft Edge、Safari、微信内置浏览器', style='List Bullet')
doc.add_paragraph('平台兼容：见数（Credamo）iframe 内嵌，高度自适应，无双层滚动条', style='List Bullet')
doc.add_paragraph('网络要求：纯前端运行，无需外部网络请求（知识库和文案均本地嵌入）', style='List Bullet')
doc.add_paragraph('移动端适配：CSS 媒体查询 @media (max-width: 480px)，自动调整字号和间距', style='List Bullet')

# ========== 保存 ==========
output_path = 'E:/大二下/厚粲杯/demo/Agent工作原理与执行流程.docx'
doc.save(output_path)
print(f'Word doc generated: {output_path}')
print(f'File size: {os.path.getsize(output_path)} bytes')
