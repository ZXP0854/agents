#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成《厚粲杯 AIGC 绿色消费实验 Agent 技术文档》Word 文档。"""

import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)

import generate_agents as GA  # noqa: E402
from docx import Document  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

OUT = os.path.join(BASE, '厚粲杯_AIGC绿色消费实验Agent技术文档.docx')

# 与 generate_agents.py 中 JS 保持一致的组别语气收尾（自然融入核心操纵元素）
CLOSING = {
    'collaboration': {
        '优点': '它（们）的这些优点，正是我们共同选择环保产品、为环保做出贡献的理由。',
        '好处': '这也让我们在为环保做出贡献的同时，收获更多。',
        '缺点': '虽然有一些不足，但为环保计，它仍值得我们一起选择。',
        '用途': '日常使用中，我们也能一起为环保尽一份力。',
        '价格': '多付一点，是我们一起为环保尽的心意。',
        '安全': '健康又环保，值得我们一起选择。',
        '耐用': '健康又环保，值得我们一起选择。',
        '对比': '综合来看，加入我们一起选择环保款，就是为环保做出贡献。',
        '推荐': '加入我们，一起为环保做出贡献。',
        'eco': '这也正是我们共同守护环境的方式。',
    },
    'non_collaboration': {
        '优点': '这些优点，正是绿色消费的优选理由。',
        '好处': '这也是您为绿色消费尽的一份力。',
        '缺点': '虽然有一些不足，但为环保计，您仍可考虑它。',
        '用途': '日常使用中，您也能为绿色消费尽一份力。',
        '价格': '多付一点，是您为环保尽的心意。',
        '安全': '健康又环保，值得您选择。',
        '耐用': '健康又环保，值得您选择。',
        '对比': '综合来看，请您选择绿色消费，环保款更值得考虑。',
        '推荐': '请您选择绿色消费，环保款值得考虑。',
        'eco': '这也正是您践行绿色消费的方式。',
    },
}

LABEL = {'优点': '优点', '好处': '好处', '缺点': '缺点', '用途': '用途',
         '价格': '价格', '安全': '安全性', '耐用': '耐用性'}


def set_font(run, name='微软雅黑', size=10.5, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)


def p(doc, text, size=10.5, bold=False, color=None, style=None):
    para = doc.add_paragraph(style=style)
    run = para.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return para


def heading(doc, text, level=1):
    h = doc.add_heading('', level=level)
    run = h.add_run(text)
    set_font(run, size={1: 16, 2: 13, 3: 11.5}.get(level, 11), bold=True,
             color=(0x1F, 0x3D, 0x24))
    return h


def build_doc():
    doc = Document()

    title = doc.add_heading('', 0)
    run = title.add_run('厚粲杯 · AIGC 绿色消费实验 Agent 技术文档')
    set_font(run, size=18, bold=True, color=(0x1F, 0x3D, 0x24))
    p(doc, '第二届"厚粲杯"全国大学生心理与认知智能测评挑战赛 参赛项目', size=11, color=(0x55, 0x55, 0x55))

    # ========== 一、技术栈 ==========
    heading(doc, '一、技术栈', 1)
    p(doc, '本套实验 Agent 为纯前端实现，无后端依赖，可直接嵌入见数（Credamo）平台运行。', size=10.5)
    for k, v in [
        ('生成脚本', 'Python 3（generate_agents.py 生成 12 个 Agent，wrap_iframe.py / wrap_noscript.py 生成见数兼容版本）'),
        ('Agent 本体', 'HTML5 + CSS3 + 原生 JavaScript（ES5，无任何第三方库，离线可运行）'),
        ('见数兼容方案 1', 'agents_iframe：内层 HTML 整体 base64 编码，嵌入 iframe 的 data URI，隐藏 <script>/<style> 标签'),
        ('见数兼容方案 2', 'agents_noscript：零 <script> 标签，JS 经 base64 编码存入隐藏 <div>，由 <body onload> 用 new Function 解码执行'),
        ('数据回传', 'window.postMessage（向宿主页面报告 session_init / user_input / product_response / qa_response / error_response 等事件）'),
        ('打包与版本管理', 'zip 打包（agents.zip）+ Git 版本管理'),
        ('本文档生成', 'python-docx'),
    ]:
        para = doc.add_paragraph()
        r1 = para.add_run('· ' + k + '：')
        set_font(r1, bold=True)
        r2 = para.add_run(v)
        set_font(r2)

    # ========== 二、实现原理 ==========
    heading(doc, '二、实现原理', 1)

    heading(doc, '2.1 实验设计', 2)
    p(doc, '采用「品类 × 组别」结构：6 个品类（笔类、包装胶带、垃圾袋、牙刷、纸杯、纸张）× 2 个组别'
           '（协作规范诉求组 / 非协作规范诉求组）= 12 个独立 Agent。每个 Agent 内置该品类的普通款与环保可降解款'
           '两类商品信息（外观描述 + 价格）。协作组与非协作组仅在干预文案的规范诉求关键词上区分，'
           '从而实现对"协作规范诉求"这一自变量的操纵。')

    heading(doc, '2.2 交互流程（分支逻辑）', 2)
    for k, v in [
        ('① 追问检测', '若输入命中追问关键词（优点/好处/缺点/用途/价格/安全/耐用/对比/推荐/环保概念），'
                    '则从 A/B 知识库中检索对应内容，并按组别语气自然收尾回答；'),
        ('② 环保款检测', '若输入命中环保款别名（如"咖啡渣环保中性笔""竹牙刷"），'
                     '提示"本次请输入普通版本商品，请重新输入对应普通商品名称。"；'),
        ('③ 普通款检测', '若输入命中本品类普通款别名，输出"本店有两种X售卖：" + A/B 外观/价格 + 干预文案；'),
        ('④ 兜底报错', '其余输入按品类固定报错（如"抱歉，暂不支持该商品，请输入与笔有关的商品。"）。'),
    ]:
        para = doc.add_paragraph()
        r1 = para.add_run(k)
        set_font(r1, bold=True)
        r2 = para.add_run(v)
        set_font(r2)

    heading(doc, '2.3 关键词匹配机制', 2)
    p(doc, '采用「别名包含匹配 + 关键词映射」两层机制：'
           '① 商品识别——将输入与普通款/环保款别名列表做不区分大小写的子串匹配；'
           '② 追问识别——将输入与追问类型关键词表逐项匹配，确定回答类型（优点/价格/对比等）；'
           '③ 目标识别——通过 A/B 指代标记（如"普通款""环保款""B的""选A"）判断被试询问的是 A、B 还是两者。')

    heading(doc, '2.4 组别操纵（核心操纵句自然融入）', 2)
    p(doc, '首次输入商品时输出完整核心操纵句：协作组以"加入我们，让我们一起为环保做出贡献！"开头；'
           '非协作组以"请选择绿色消费！"开头。在后续追问回答中，核心操纵元素（协作组：加入我们/一起/共同/我们；'
           '非协作组：请选择/您/个人/绿色消费）以自然语句收尾的方式融入，避免生硬重复整句，'
           '从而在多次交互中持续、自然地进行规范诉求操纵。')

    heading(doc, '2.5 见数兼容方案', 2)
    p(doc, '见数平台对问卷中嵌入的 HTML 有一定限制，本项目提供三种形态：'
           '① agents/ 标准版（含 <script>，供本地直接打开测试）；'
           '② agents_iframe/ 将内层 HTML 整体 base64 编码后嵌入 iframe，隐藏 script/style 标签；'
           '③ agents_noscript/ 将 JS 编码后经 <body onload> 执行，彻底移除 <script> 标签。'
           '三种形态的交互逻辑完全一致。')

    # ========== 三、词库 ==========
    heading(doc, '三、词库', 1)

    heading(doc, '3.1 品类与商品词库（普通款 / 环保款）', 2)
    for cat in GA.CATEGORIES:
        p(doc, '【%s】普通款：%s（%s元）；环保款：%s（%s元）' % (
            cat['category'], cat['normal_name'], cat['normal_price'],
            cat['eco_name'], cat['eco_price']), bold=True)
        p(doc, '普通款别名：' + '、'.join(cat['normal_aliases']))
        p(doc, '环保款别名：' + '、'.join(cat['eco_aliases']))
        p(doc, '普通款外观：' + cat['normal_appearance'])
        p(doc, '环保款外观：' + cat['eco_appearance'])

    heading(doc, '3.2 组别核心操纵句与关键词', 2)
    for k, v in [
        ('协作组', '核心操纵句「加入我们，让我们一起为环保做出贡献！」；关键词：加入我们 / 一起 / 共同 / 我们'),
        ('非协作组', '核心操纵句「请选择绿色消费！」；关键词：请选择 / 您 / 个人 / 绿色消费'),
    ]:
        para = doc.add_paragraph()
        r1 = para.add_run('· ' + k + '：')
        set_font(r1, bold=True)
        r2 = para.add_run(v)
        set_font(r2)

    heading(doc, '3.3 追问类型关键词表', 2)
    qt_keys = {
        '优点': ['优点', '优势', '特色', '亮点', '哪里好', '好在', '好不好', '怎么样'],
        '好处': ['好处', '有什么用', '有什么好', '作用', '帮助'],
        '缺点': ['缺点', '不足', '劣势', '缺陷', '坏处', '不好', '问题'],
        '用途': ['用途', '怎么用', '使用场景', '能做什么', '适用', '适合', '场景', '干什么用'],
        '价格': ['价格', '多少钱', '贵', '便宜', '成本', '花费', '划算', '价位'],
        '安全': ['安全', '有毒', '无害', '健康', '放心', '标准', '检测', '危害'],
        '耐用': ['耐用', '质量', '寿命', '持久', '容易坏', '结实', '保存', '保质', '储存', '好用吗'],
        '对比': ['对比', '区别', '比较', '相比', 'vs', '哪个好', '有什么不同', '差别', '不一样'],
        '推荐': ['推荐', '建议', '买哪个', '选哪个', '应该买', '值得买', '买什么', '怎么选'],
    }
    for k, v in qt_keys.items():
        para = doc.add_paragraph()
        r1 = para.add_run('· ' + k + '：')
        set_font(r1, bold=True)
        r2 = para.add_run('、'.join(v))
        set_font(r2)

    heading(doc, '3.4 A/B 知识库（每品类 7 项 + 对比 + 推荐）', 2)
    for cat in GA.CATEGORIES:
        p(doc, '【%s】' % cat['category'], bold=True, size=11)
        for item in ['优点', '好处', '缺点', '用途', '价格', '安全', '耐用']:
            p(doc, 'A（%s）%s：%s' % (cat['normal_name'], item, cat['normal_kb'][item]))
            p(doc, 'B（%s）%s：%s' % (cat['eco_name'], item, cat['eco_kb'][item]))
        p(doc, '对比：' + cat['compare'])
        p(doc, '推荐：' + cat['recommend'])

    heading(doc, '3.5 组别语气收尾词库（自然融入）', 2)
    for gname, gmap in [('协作组', CLOSING['collaboration']), ('非协作组', CLOSING['non_collaboration'])]:
        p(doc, '【%s】' % gname, bold=True)
        for k in ['优点', '好处', '缺点', '用途', '价格', '安全', '对比', '推荐', 'eco']:
            p(doc, '· %s：%s' % (k, gmap[k]))

    # ========== 四、对应回答 ==========
    heading(doc, '四、对应回答', 1)

    heading(doc, '4.1 固定话术', 2)
    p(doc, '欢迎语：顾客您好！欢迎光临小林超市，我是本店的 AI 智能助手小林，请输入您的商品信息。')
    p(doc, '环保款提示：本次请输入普通版本商品，请重新输入对应普通商品名称。')
    p(doc, '报错（按品类）：抱歉，暂不支持该商品，请输入与X有关的商品。（X 为该品类量词，如"笔""牙刷"）')
    p(doc, '输入框灰字：初始「请输入产品名称…」；输入正确商品后变为「您现在可以询问两款产品的具体信息」。')

    heading(doc, '4.2 普通商品回复（输入本品类普通款）', 2)
    p(doc, '回复格式：「本店有两种X售卖：」+ A（普通款）名称/外观/价格 + B（环保款）名称/外观/价格 + 干预文案。')
    for group_type, group_label in [('collaboration', '协作组'), ('non_collaboration', '非协作组')]:
        copy_dict = GA.COLLABORATION_COPY if group_type == 'collaboration' else GA.NON_COLLABORATION_COPY
        for cat in GA.CATEGORIES:
            resp = GA.build_normal_response(cat, copy_dict[cat['copy_key']])
            p(doc, '【%s · %s】%s' % (cat['category'], group_label, resp), size=9.5)

    heading(doc, '4.3 追问回答（示例：优点 / 价格 / 对比 / 推荐）', 2)
    p(doc, '追问回答 = A/B 知识库内容 + 组别语气收尾。下面以「笔类」为例给出协作组与非协作组的组装结果：')
    for group_type, group_label, gkey in [('collaboration', '协作组', 'collaboration'), ('non_collaboration', '非协作组', 'non_collaboration')]:
        cat = GA.CATEGORIES[0]  # 笔类
        cl = CLOSING[gkey]
        p(doc, '【%s · 笔类】' % group_label, bold=True)
        p(doc, '问「优点」：A（%s）的优点：%s%s' % (cat['normal_name'], cat['normal_kb']['优点'], cl['优点']))
        p(doc, '问「价格」：A（%s）的价格：%s B（%s）的价格：%s%s' % (
            cat['normal_name'], cat['normal_kb']['价格'], cat['eco_name'], cat['eco_kb']['价格'], cl['价格']))
        p(doc, '问「对比」：%s%s' % (cat['compare'], cl['对比']))
        p(doc, '问「推荐」：%s%s' % (cat['recommend'], cl['推荐']))

    p(doc, '（其余品类按同一模式组装，A/B 知识与收尾语见第三节词库。）', size=9.5, color=(0x66, 0x66, 0x66))

    doc.save(OUT)
    return OUT


if __name__ == '__main__':
    out = build_doc()
    print('文档已生成：%s' % out)
    print('大小：%d 字节' % os.path.getsize(out))
